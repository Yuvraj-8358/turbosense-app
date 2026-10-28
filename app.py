import io
import json
import math
import os
import threading
from functools import lru_cache
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import pandas as pd
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

ROOT = Path(__file__).resolve().parent
MODEL_DIR = Path(os.getenv("MODEL_DIR", ROOT / "models"))
MANIFEST = json.loads((ROOT / "manifest.json").read_text(encoding="utf-8"))
DATASETS = set(MANIFEST["datasets"])
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UPLOAD_ROWS = 200
MODEL_LOAD_LOCK = threading.Lock()

app = FastAPI(
    title="TURBOSENSE",
    version=MANIFEST.get("version", "1.0.0"),
    description="NASA C-MAPSS RUL and health-state inference",
)
app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")


class PredictionRequest(BaseModel):
    subset: str = Field(description="One of FD001, FD002, FD003, or FD004")
    rows: list[dict[str, Any]] = Field(
        min_length=1,
        max_length=MAX_UPLOAD_ROWS,
        description="Engine-cycle rows using the feature names for the selected subset",
    )


def features_for(subset: str) -> list[str]:
    model_info = MANIFEST.get("selected_models", {}).get(subset, {})
    features = model_info.get("features")
    return list(features or MANIFEST.get("features", []))


@lru_cache(maxsize=1)
def load_bundle(subset: str):
    with MODEL_LOAD_LOCK:
        path = MODEL_DIR / f"{subset}_model.joblib"
        if not path.is_file():
            raise FileNotFoundError(f"Model artifact not found: {path.name}")
        bundle = joblib.load(path)
        required = {"regressor", "classifier", "features", "rul_cap"}
        if not required.issubset(bundle):
            missing = sorted(required - set(bundle))
            raise RuntimeError(f"{path.name} is not a combined TURBOSENSE artifact; missing keys: {missing}")
        return bundle


def docx_to_frame(content: bytes) -> pd.DataFrame:
    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="Could not read this Word file. Save it as .docx and try again.",
        ) from exc

    for table in document.tables:
        matrix = [
            [cell.text.strip() for cell in row.cells]
            for row in table.rows
        ]
        matrix = [row for row in matrix if any(cell for cell in row)]
        if len(matrix) < 2:
            continue
        headers = matrix[0]
        if not headers or any(not header for header in headers):
            continue
        records = []
        for row in matrix[1:]:
            values = (row + [""] * len(headers))[:len(headers)]
            if any(value for value in values):
                records.append(values)
        if records:
            return pd.DataFrame(records, columns=headers)

    supported_fields = set(MANIFEST.get("features", [])) | {"unit", "cycle"}
    pairs: dict[str, str] = {}
    for paragraph in document.paragraphs:
        line = paragraph.text.strip()
        delimiter = ":" if ":" in line else "=" if "=" in line else None
        if not delimiter:
            continue
        key, value = line.split(delimiter, 1)
        key = key.strip()
        if key in supported_fields and value.strip():
            pairs[key] = value.strip()
    if pairs:
        return pd.DataFrame([pairs])

    raise HTTPException(
        status_code=422,
        detail="The Word document needs a table with sensor names in its first row and readings below.",
    )


def json_safe(value: Any) -> Any:
    if pd.isna(value):
        return None
    if isinstance(value, np.generic):
        return value.item()
    return value


def uploaded_frame(filename: str, content: bytes) -> pd.DataFrame:
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".csv":
            return pd.read_csv(io.BytesIO(content), encoding="utf-8-sig", skipinitialspace=True)
        if extension == ".docx":
            return docx_to_frame(content)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="Could not parse this file. Check the CSV headers or the Word table format.",
        ) from exc
    raise HTTPException(
        status_code=415,
        detail="Upload a CSV file or a modern Word document saved as .docx.",
    )


@app.get("/", include_in_schema=False)
def home():
    return FileResponse(ROOT / "static" / "index.html")


@app.get("/healthz")
def healthz():
    missing = [
        f"{dataset}_model.joblib"
        for dataset in sorted(DATASETS)
        if not (MODEL_DIR / f"{dataset}_model.joblib").is_file()
    ]
    return {
        "status": "ok" if not missing else "models_missing",
        "datasets": sorted(DATASETS),
        "missing_artifacts": missing,
    }


@app.get("/api/manifest")
def get_manifest():
    return MANIFEST


@app.post("/api/parse-upload")
async def parse_upload(
    subset: str = Form(...),
    file: UploadFile = File(...),
):
    subset = subset.strip().upper()
    if subset not in DATASETS:
        raise HTTPException(status_code=422, detail="Choose FD001, FD002, FD003, or FD004.")

    filename = Path(file.filename or "").name
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".docx"}:
        raise HTTPException(
            status_code=415,
            detail="Upload a CSV file or a modern Word document saved as .docx.",
        )

    content = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if not content:
        raise HTTPException(status_code=422, detail="The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Files must be 5 MB or smaller.")

    frame = uploaded_frame(filename, content)
    if frame.empty:
        raise HTTPException(status_code=422, detail="The uploaded file contains no sensor readings.")

    frame.columns = [str(column).strip() for column in frame.columns]
    if any(not column for column in frame.columns):
        raise HTTPException(status_code=422, detail="Every input column needs a sensor-name header.")
    if frame.columns.duplicated().any():
        raise HTTPException(status_code=422, detail="The file contains duplicate column names.")

    required_features = features_for(subset)
    missing = [feature for feature in required_features if feature not in frame.columns]
    if missing:
        raise HTTPException(
            status_code=422,
            detail={
                "message": "The file is missing required inputs for " + subset + ".",
                "missing_features": missing,
            },
        )

    if len(frame) > MAX_UPLOAD_ROWS:
        raise HTTPException(
            status_code=413,
            detail=f"Use no more than {MAX_UPLOAD_ROWS} readings in one upload.",
        )

    for feature in required_features:
        numeric = pd.to_numeric(frame[feature], errors="coerce")
        values = numeric.to_numpy(dtype=float)
        if not np.isfinite(values).all():
            raise HTTPException(
                status_code=422,
                detail=f"Column {feature} must contain finite numbers in every row.",
            )
        frame[feature] = numeric

    if "cycle" in frame.columns:
        cycles = pd.to_numeric(frame["cycle"], errors="coerce")
        if not np.isfinite(cycles.to_numpy(dtype=float)).all():
            raise HTTPException(status_code=422, detail="Cycle values must be finite numbers.")
        frame["cycle"] = cycles

    if "unit" in frame.columns:
        frame["unit"] = frame["unit"].fillna("").astype(str).str.strip()

    rows = [
        {column: json_safe(value) for column, value in record.items()}
        for record in frame.to_dict(orient="records")
    ]
    return {
        "subset": subset,
        "filename": filename,
        "row_count": len(rows),
        "columns": list(frame.columns),
        "rows": rows,
    }


@app.post("/api/predict")
def predict(request: PredictionRequest):
    subset = request.subset.strip().upper()
    if subset not in DATASETS:
        raise HTTPException(status_code=422, detail="subset must be FD001, FD002, FD003, or FD004")
    try:
        bundle = load_bundle(subset)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except MemoryError as exc:
        raise HTTPException(
            status_code=503,
            detail=f"{subset} is temporarily unavailable because the model exceeded the service memory limit.",
        ) from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Could not load {subset} model bundle: {exc}") from exc

    frame = pd.DataFrame(request.rows)
    missing = [feature for feature in bundle["features"] if feature not in frame.columns]
    if missing:
        raise HTTPException(
            status_code=422,
            detail={"message": "Input is missing model features", "missing_features": missing},
        )

    x = frame[bundle["features"]].apply(pd.to_numeric, errors="coerce")
    if x.isna().any().any() or not np.isfinite(x.to_numpy(dtype=float)).all():
        raise HTTPException(status_code=422, detail="Feature values must be finite numbers")

    if "unit" in frame.columns and "cycle" in frame.columns:
        cycles = pd.to_numeric(frame["cycle"], errors="coerce")
        if cycles.isna().any() or not np.isfinite(cycles.to_numpy(dtype=float)).all():
            raise HTTPException(status_code=422, detail="cycle values must be finite numbers when unit and cycle are provided")
        frame["cycle"] = cycles
        frame = frame.sort_values("cycle").groupby("unit", sort=False).tail(1)
        x = frame[bundle["features"]].apply(pd.to_numeric, errors="coerce")

    try:
        rul = np.clip(bundle["regressor"].predict(x), 0, bundle["rul_cap"])
        labels = bundle["classifier"].predict(x)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Prediction failed: {exc}") from exc

    predictions = [
        {"predicted_rul": round(float(rul_value), 2), "health": str(label)}
        for rul_value, label in zip(rul, labels)
    ]
    return {"subset": subset, "prediction_count": len(predictions), "predictions": predictions}
